from flask import Flask, render_template, request
import fitz
from docx import Document
import os

app = Flask(__name__)

# ==============================
# Upload Folder
# ==============================

UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


# ==============================
# Job Required Skills
# ==============================

JOB_SKILLS = {

    "Data Analyst": [
        "python",
        "sql",
        "excel",
        "power bi",
        "statistics",
        "pandas",
        "numpy",
        "data visualization"
    ],

    "Data Scientist": [
        "python",
        "machine learning",
        "statistics",
        "pandas",
        "numpy",
        "scikit-learn",
        "sql",
        "deep learning"
    ],

    "Frontend Developer": [
        "html",
        "css",
        "javascript",
        "react",
        "git",
        "responsive design"
    ],

    "Backend Developer": [
        "python",
        "java",
        "node.js",
        "sql",
        "mongodb",
        "api",
        "git"
    ]
}


# ==============================
# Extract Text From PDF
# ==============================

def extract_pdf_text(file_path):

    text = ""

    try:
        document = fitz.open(file_path)

        for page in document:
            text += page.get_text("text") + "\n"

        document.close()

    except Exception as e:
        print("PDF extraction error:", e)

    return text


# ==============================
# Extract Text From DOCX
# ==============================

def extract_docx_text(file_path):

    text = ""

    try:
        document = Document(file_path)

        # Paragraphs
        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"

        # Tables
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + "\n"

    except Exception as e:
        print("DOCX extraction error:", e)

    return text


# ==============================
# Extract Resume Text
# ==============================

def extract_text(file_path):

    extension = file_path.lower().split(".")[-1]

    if extension == "pdf":
        return extract_pdf_text(file_path)

    elif extension == "docx":
        return extract_docx_text(file_path)

    else:
        return ""


# ==============================
# Find Skills
# ==============================

def find_skills(text, skills):

    # Convert complete resume text to lowercase
    text = text.lower()

    found = []

    for skill in skills:

        skill_lower = skill.lower()

        if skill_lower in text:
            found.append(skill)

    return found


# ==============================
# Home Route
# ==============================

@app.route("/", methods=["GET", "POST"])
def home():

    if request.method == "POST":

        # Get uploaded resume
        resume = request.files.get("resume")

        # Get selected job
        job_role = request.form.get("job_role")

        # Check resume
        if not resume or resume.filename == "":
            return render_template(
                "index.html",
                error="Please upload your resume."
            )

        # Check job
        if job_role not in JOB_SKILLS:
            return render_template(
                "index.html",
                error="Please select a valid job role."
            )

        # ==============================
        # Save Resume
        # ==============================

        filename = resume.filename

        file_path = os.path.join(
            app.config["UPLOAD_FOLDER"],
            filename
        )

        resume.save(file_path)

        # ==============================
        # Extract Resume Text
        # ==============================

        resume_text = extract_text(file_path)

        # Debugging
        print("\n===================================")
        print("EXTRACTED RESUME TEXT")
        print("===================================")
        print(resume_text)
        print("===================================\n")

        # Check if text was extracted
        if not resume_text.strip():

            return render_template(
                "index.html",
                error=(
                    "Could not extract text from the resume. "
                    "Please upload a text-based PDF or DOCX file."
                )
            )

        # ==============================
        # Required Skills
        # ==============================

        required_skills = JOB_SKILLS[job_role]

        # ==============================
        # Find Skills
        # ==============================

        found_skills = find_skills(
            resume_text,
            required_skills
        )

        # ==============================
        # Missing Skills
        # ==============================

        missing_skills = [
            skill
            for skill in required_skills
            if skill not in found_skills
        ]

        # ==============================
        # Calculate Score
        # ==============================

        total_skills = len(required_skills)

        if total_skills > 0:
            score = round(
                (len(found_skills) / total_skills) * 100
            )
        else:
            score = 0

        # ==============================
        # Debugging Skill Result
        # ==============================

        print("JOB ROLE:", job_role)
        print("REQUIRED SKILLS:", required_skills)
        print("FOUND SKILLS:", found_skills)
        print("MISSING SKILLS:", missing_skills)
        print("MATCH SCORE:", score)

        # ==============================
        # Send Result To HTML
        # ==============================

        return render_template(
            "index.html",
            result=True,
            score=score,
            job_role=job_role,
            found_skills=found_skills,
            missing_skills=missing_skills
        )

    return render_template("index.html")


# ==============================
# Run Flask Application
# ==============================

if __name__ == "__main__":
    app.run(
        debug=True,
        host="127.0.0.1",
        port=5000
    )
